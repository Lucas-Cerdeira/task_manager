from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from io import BytesIO
from typing import List
import os
import tempfile
from app.schemas.document import SenderRecipientPair, Address


class DocumentService:
    
    @staticmethod
    def create_sender_recipient_document(pairs: List[SenderRecipientPair]) -> str:
        """Create a Word document with sender and recipient pairs, one pair per page."""
        
        doc = Document()
        
        for i, pair in enumerate(pairs):
            if i > 0:
                # Add page break before each new pair (except the first one)
                doc.add_page_break()
            
            # Add sender section
            sender_title = doc.add_paragraph("REMETENTE")
            sender_title.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            sender_title.runs[0].bold = True
            
            DocumentService._add_address_info(doc, pair.sender)
            
            # Add some space between sender and recipient
            doc.add_paragraph()
            
            # Add recipient section  
            recipient_title = doc.add_paragraph("DESTINATÁRIO")
            recipient_title.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            recipient_title.runs[0].bold = True
            
            DocumentService._add_address_info(doc, pair.recipient)
        
        # Save document to a temporary file
        temp_dir = tempfile.gettempdir()
        filename = f"sender_recipient_document.docx"
        filepath = os.path.join(temp_dir, filename)
        
        doc.save(filepath)
        
        return filepath
    
    @staticmethod
    def _add_address_info(doc: Document, address: Address):
        """Add address information to the document."""
        doc.add_paragraph(address.name)
        
        # Build address line
        address_parts = []
        if address.street:
            street_info = address.street
            if address.number:
                street_info += f" {address.number}"
            address_parts.append(street_info)
        
        if address.neighborhood:
            address_parts.append(address.neighborhood)
        
        if address_parts:
            doc.add_paragraph(" ".join(address_parts))
        
        # Add postal code
        if address.postal_code:
            doc.add_paragraph(f"CEP {address.postal_code}")
        
        # Add city and state
        location_parts = []
        if address.city:
            location_parts.append(address.city)
        if address.state:
            location_parts.append(address.state)
        
        if location_parts:
            doc.add_paragraph(" - ".join(location_parts))